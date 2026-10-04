#!/usr/bin/env python3
"""Build the TikTok content-plan tracker (.xlsx) from docs/TIKTOK_CONTENT_PLAN_M1.md.

Usage: python3 tools/build_plan_sheet.py out.xlsx
The .xlsx is uploaded to Google Drive and converted to Google Sheets.
"""
import re
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
PLAN = ROOT / "docs" / "TIKTOK_CONTENT_PLAN_M1.md"
REPO = "https://github.com/svistsvat1-debug/work-cloud-/tree/workproces/"

STYLE_NAMES = {
    "DN": "Dark Neon", "YP": "Yellow Punch", "CL": "Clean Light", "PN": "Phone Native",
    "TM": "Terminal", "RW": "Retro Web", "BP": "Blueprint",
}
STYLE_ROWS = [
    ("DN", "Dark Neon", "Темний фон #0A0A0A, неоново-жовті акценти, glow. Основний стиль (як #01)."),
    ("YP", "Yellow Punch", "Повністю жовтий фон, чорний текст і UI."),
    ("CL", "Clean Light", "Світлий фон #F5F5F3, чорний текст, жовтий маркер-виділення."),
    ("PN", "Phone Native", "Екран телефону: сповіщення, переписка в DM, Notes."),
    ("TM", "Terminal", "Чорний фон, моноширинний шрифт, логи автоматизації."),
    ("RW", "Retro Web", "Сайт у стилі 2005–2014: bevel-кнопки, Times, лічильник відвідувань."),
    ("BP", "Blueprint", "Темно-синя сітка, білі лінії-схеми."),
]

# Production facts for finished videos (filled after each render).
DONE = {
    1: {
        "summary": "AI chat → AI names 3 competitors, not you → AI reads websites → vague site gets flagged → fix: clear services, city, FAQ → your business lands #1",
        "status": "Готово",
        "length": 30,
        "files": REPO + "videos/01-ask-chatgpt",
        "caption": "Ask AI who to hire in your city. Is your business on the list? 👀 #smallbusiness #websitetips #aiforbusiness #chatgpt #localbusiness",
        "notes": "Голос Kokoro (тимчасово, до ElevenLabs). Є версія без музики.",
    },
}

HEADERS = [
    "#", "День", "Дата", "Пілар", "Стиль", "Формат", "Хук (перші 2 с)", "Тема й кут", "Сценарій коротко",
    "CTA", "План тривалості", "Статус", "Факт, с", "Файли", "Caption і хештеги", "Опубліковано", "Посилання TikTok",
    "Перегляди", "Лайки", "Коментарі", "Поширення", "Збереження", "Нові підписники", "Переходи на сайт",
    "Engagement rate", "Нотатки",
]
WIDTHS = [5, 7, 11, 15, 14, 18, 34, 34, 46, 18, 9, 13, 8, 22, 40, 13, 22, 11, 9, 11, 11, 12, 12, 12, 12, 34]
STATUSES = ["Заплановано", "В роботі", "Готово", "Опубліковано"]

HEAD_ROW = 5  # header row; data starts below
FONT = "Arial"


def clean(cell):
    return re.sub(r"[*`]", "", cell).strip()


def parse_plan():
    rows = []
    for line in PLAN.read_text().splitlines():
        if not re.match(r"^\| \d+ \|", line):
            continue
        c = [x.strip() for x in line.strip().strip("|").split("|")]
        num, day, pillar, style_fmt, hook, topic, summary, cta, length = c[:9]
        code, _, fmt = style_fmt.partition("·")
        rows.append({
            "num": int(num), "day": int(day), "pillar": pillar, "style": code.strip(), "format": fmt.strip(),
            "hook": re.sub(r"^'(.*)'$", "\u201c\\1\u201d", clean(hook).strip('"')), "topic": clean(topic), "summary": clean(summary),
            "cta": "Сайт (link in bio)" if "SITE" in cta else "Follow for more",
            "length": clean(length),
        })
    assert len(rows) == 25, f"expected 25 rows, got {len(rows)}"
    return rows


def style_label(code):
    parts = [STYLE_NAMES.get(p.strip(), p.strip()) for p in code.replace("→", ">").split(">")]
    return " → ".join(parts)


def write_csv(out):
    """Plain CSV of the same tracker (Drive converts it to a Google Sheet)."""
    import csv
    rows = [
        ["TikTok контент-план, місяць 1 (25 відео)"],
        ["Старт (День 1)", "", "2026-10-04", "Дата старту, від неї рахуються всі дати."],
        [],  # summary row: formulas are added through the Sheets API (locale-independent syntax)
        [],
        HEADERS,
    ]
    for r, p in enumerate(parse_plan(), HEAD_ROW + 1):
        done = DONE.get(p["num"], {})
        rows.append([
            p["num"], p["day"], f"=$C$2+B{r}-1", p["pillar"], style_label(p["style"]), p["format"],
            p["hook"], p["topic"], done.get("summary", p["summary"]), p["cta"], p["length"],
            done.get("status", "Заплановано"), done.get("length", ""), done.get("files", ""), done.get("caption", ""),
            "", "", "", "", "", "", "", "", "", "", done.get("notes", ""),
        ])
    with open(out, "w", newline="") as fh:
        csv.writer(fh).writerows(rows)
    print(f"saved {out}")


def main(out):
    if out.endswith(".csv"):
        return write_csv(out)
    wb = Workbook()
    ws = wb.active
    ws.title = "Контент-план"
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    base = Font(name=FONT, size=10)

    ws["A1"] = "TikTok контент-план, місяць 1 (25 відео)"
    ws["A1"].font = Font(name=FONT, size=14, bold=True)
    ws["A2"] = "Старт (День 1)"
    ws["C2"] = date(2026, 10, 4)
    ws["C2"].number_format = "dd.mm.yyyy"
    ws["C2"].fill = PatternFill("solid", fgColor="FFF2A8")
    ws["D2"] = "Жовта клітинка: дата старту, від неї рахуються всі дати."
    ws["A3"] = "Готово"
    ws["C3"] = f'=COUNTIF(L{HEAD_ROW + 1}:L1000,"Готово")+COUNTIF(L{HEAD_ROW + 1}:L1000,"Опубліковано")'
    ws["D3"] = "Опубліковано"
    ws["F3"] = f'=COUNTIF(L{HEAD_ROW + 1}:L1000,"Опубліковано")'
    ws["G3"] = "Відео з CTA на сайт"
    ws["H3"] = f'=COUNTIF(J{HEAD_ROW + 1}:J1000,"Сайт (link in bio)")'
    for ref in ("A2", "A3", "D3", "G3"):
        ws[ref].font = Font(name=FONT, size=10, bold=True)
    for ref in ("C2", "C3", "F3", "H3", "D2"):
        ws[ref].font = base
    ws["D2"].font = Font(name=FONT, size=9, italic=True, color="666666")

    for i, h in enumerate(HEADERS, 1):
        c = ws.cell(row=HEAD_ROW, column=i, value=h)
        c.font = Font(name=FONT, size=10, bold=True, color="FFE600")
        c.fill = PatternFill("solid", fgColor="111111")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    ws.row_dimensions[HEAD_ROW].height = 32

    for r, p in enumerate(parse_plan(), HEAD_ROW + 1):
        done = DONE.get(p["num"], {})
        vals = [
            p["num"], p["day"], f"=$C$2+B{r}-1", p["pillar"], style_label(p["style"]), p["format"],
            p["hook"], p["topic"], done.get("summary", p["summary"]), p["cta"], p["length"],
            done.get("status", "Заплановано"), done.get("length"), done.get("files"), done.get("caption"),
            None, None, None, None, None, None, None, None, None,
            f'=IF(OR(R{r}="",R{r}=0),"",(S{r}+T{r}+U{r}+V{r})/R{r})', done.get("notes"),
        ]
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=col, value=v)
            c.font = base
            c.border = border
            c.alignment = Alignment(vertical="top", wrap_text=col in (6, 7, 8, 9, 14, 15, 26))
        ws.cell(row=r, column=3).number_format = "dd.mm"
        ws.cell(row=r, column=25).number_format = "0.0%"
        if done.get("files"):
            ws.cell(row=r, column=14).hyperlink = done["files"]
            ws.cell(row=r, column=14).value = "videos/" + done["files"].rsplit("/", 1)[-1]
            ws.cell(row=r, column=14).font = Font(name=FONT, size=10, color="1155CC", underline="single")
        ws.cell(row=r, column=7).font = Font(name=FONT, size=10, bold=True)

    last = HEAD_ROW + 25
    data = f"A{HEAD_ROW + 1}:Z{last}"
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"L{HEAD_ROW + 1}:L{last + 50}")
    first = HEAD_ROW + 1
    ws.conditional_formatting.add(data, FormulaRule(formula=[f'$L{first}="Опубліковано"'], fill=PatternFill("solid", fgColor="D9EAF7")))
    ws.conditional_formatting.add(data, FormulaRule(formula=[f'$L{first}="Готово"'], fill=PatternFill("solid", fgColor="DDF2D8")))
    ws.conditional_formatting.add(data, FormulaRule(formula=[f'$L{first}="В роботі"'], fill=PatternFill("solid", fgColor="FFF6CC")))
    ws.conditional_formatting.add(f"J{first}:J{last}", FormulaRule(formula=[f'$J{first}="Сайт (link in bio)"'], fill=PatternFill("solid", fgColor="FFE600"), font=Font(bold=True)))

    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=HEAD_ROW + 1, column=2)

    st = wb.create_sheet("Стилі")
    for i, h in enumerate(["Код", "Стиль", "Як виглядає", "Кількість відео"], 1):
        c = st.cell(row=1, column=i, value=h)
        c.font = Font(name=FONT, size=10, bold=True, color="FFE600")
        c.fill = PatternFill("solid", fgColor="111111")
        c.border = border
    for r, (code, name, desc) in enumerate(STYLE_ROWS, 2):
        vals = [code, name, desc, f"=COUNTIF('Контент-план'!E{first}:E1000,\"*\"&B{r}&\"*\")"]
        for col, v in enumerate(vals, 1):
            c = st.cell(row=r, column=col, value=v)
            c.font = base
            c.border = border
            c.alignment = Alignment(vertical="top", wrap_text=col == 3)
    for col, w in zip("ABCD", (8, 16, 60, 16)):
        st.column_dimensions[col].width = w
    st.freeze_panes = "A2"

    wb.save(out)
    print(f"saved {out}")


if __name__ == "__main__":
    main(sys.argv[1])
